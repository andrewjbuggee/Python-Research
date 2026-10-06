"""Season definitions, mean ± standard deviation statistics, and Kavin's table values.

Seasons follow the column headers of the EPCAPE seasonal-averages table
(sheet "Cloud&MetQuantities"):

    EPCAPE Avg (F23-F24)  2023-02-15 .. 2024-02-14  (whole campaign)
    Spring(AMJn23)        2023-04-01 .. 2023-06-30
    Summer(JlAS23)        2023-07-01 .. 2023-09-30
    Fall(OND23)           2023-10-01 .. 2023-12-31
    Winter(FM23+JF24)     2023-02-15 .. 2023-03-31  and  2024-01-01 .. 2024-02-14

Boundaries are applied in UTC (local PST/PDT would move 7-8 h of data at
each boundary, negligible for seasonal statistics). Every statistic is the
arithmetic mean ± sample standard deviation (ddof = 1) over all samples in
the season, so the campaign value is sample-weighted, not an average of the
four seasonal means. The sample (5-min mean, radiosonde launch, rain event,
...) is stated wherever a statistic is computed, because the standard
deviation depends on it.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

# Season key -> list of inclusive (first day, last day) ranges, UTC.
SEASONS: Dict[str, List[Tuple[str, str]]] = {
    "EPCAPE": [("2023-02-15", "2024-02-14")],
    "Spring": [("2023-04-01", "2023-06-30")],
    "Summer": [("2023-07-01", "2023-09-30")],
    "Fall": [("2023-10-01", "2023-12-31")],
    "Winter": [("2023-02-15", "2023-03-31"), ("2024-01-01", "2024-02-14")],
}
SEASON_ORDER = list(SEASONS)
# Spreadsheet column letter of each season (row 2 holds the headers).
SEASON_COLUMNS = {"EPCAPE": "E", "Spring": "F", "Summer": "G", "Fall": "H", "Winter": "I"}


# ---------------------------------------------------------------------------
# seasons and statistics
# ---------------------------------------------------------------------------
def season_mask(times, season: str) -> np.ndarray:
    """True where `times` (anything pandas can convert) fall in `season`."""
    t = pd.DatetimeIndex(pd.to_datetime(np.asarray(times)))
    if t.tz is not None:
        t = t.tz_convert("UTC").tz_localize(None)
    mask = np.zeros(len(t), dtype=bool)
    for first, last in SEASONS[season]:
        t0 = pd.Timestamp(first)
        t1 = pd.Timestamp(last) + pd.Timedelta(days=1)  # whole last day included
        mask |= (t >= t0) & (t < t1)
    return mask


def seasonal_stats(values: pd.Series, *, ddof: int = 1) -> pd.DataFrame:
    """Mean, standard deviation, median and sample count per season.

    `values` is a Series indexed by time; NaN samples are ignored. Rows are
    SEASON_ORDER; columns mean, std, median, n."""
    values = pd.Series(values).astype(float)
    rows = {}
    for season in SEASON_ORDER:
        v = values.values[season_mask(values.index, season)]
        v = v[np.isfinite(v)]
        rows[season] = {
            "mean": v.mean() if v.size else np.nan,
            "std": v.std(ddof=ddof) if v.size > ddof else np.nan,
            "median": np.median(v) if v.size else np.nan,
            "n": int(v.size),
        }
    return pd.DataFrame.from_dict(rows, orient="index")[["mean", "std", "median", "n"]]


def monthly_means(values: pd.Series) -> pd.Series:
    """Calendar-month means of the finite samples, indexed by each month's first valid sample.

    Indexing by the first sample (not the 1st of the month) keeps February
    2023, whose data start on the 15th, inside the campaign window.
    ``seasonal_stats(monthly_means(x))`` then gives mean ± std ACROSS MONTHS,
    which is much smaller than the spread of the samples themselves."""
    v = pd.Series(values).astype(float).dropna()
    if v.empty:
        return v
    groups = v.groupby(v.index.to_period("M"))
    first = groups.apply(lambda x: x.index[0])
    return pd.Series(groups.mean().values, index=pd.DatetimeIndex(first.values), name=v.name)


def seasonal_sums(values: pd.Series) -> pd.Series:
    """Sum per season (e.g. sampling hours). NaN counts as 0."""
    values = pd.Series(values).astype(float)
    return pd.Series(
        {s: float(np.nansum(values.values[season_mask(values.index, s)])) for s in SEASON_ORDER}
    )


# ---------------------------------------------------------------------------
# reading Kavin's values from the spreadsheet
# ---------------------------------------------------------------------------
# One number, optionally written as "<mantissa> * 10^<exponent>" (e.g. "1.0 * 10^2").
_NUMBER = re.compile(r"^\s*(\d+\.?\d*|\.\d+)\.?\s*(?:[*x×]\s*10\s*\^\s*([-+]?\d+))?\s*$")


def _to_float(token: str) -> float:
    m = _NUMBER.match(token)
    if not m:
        return np.nan
    value = float(m.group(1))
    if m.group(2):
        value *= 10.0 ** int(m.group(2))
    return value


def parse_cell(value) -> Tuple[float, float, str]:
    """(mean, std, note) from a table cell such as "0.58 ± 0.29".

    Handles the spellings found in the sheet: plain numbers, "a ± b",
    "a ± n/a", "1.0 * 10^2 ± 66", "5.0*10^2", ".39 ± 0.19", "0.11. ± 0.10",
    and "626 + 297" (read as ±, with a note). Returns NaN for what cannot be
    read; `note` says why or what was assumed."""
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return np.nan, np.nan, "empty"
    if isinstance(value, (int, float, np.number)):
        return float(value), np.nan, "no ± given"
    text = str(value).strip()
    if text.lower() in ("", "n/a", "na"):
        return np.nan, np.nan, "n/a"
    if text.startswith("="):
        return np.nan, np.nan, f"formula {text}"
    note = ""
    if "±" in text:
        left, right = text.split("±", 1)
    elif re.search(r"\d\s+\+\s+\.?\d", text):
        left, right = re.split(r"\s+\+\s+", text, maxsplit=1)
        note = "'+' read as '±'"
    else:
        left, right = text, ""
    mean = _to_float(left)
    right = right.strip()
    std = np.nan if right.lower() in ("", "n/a", "na") else _to_float(right)
    if right.lower() in ("n/a", "na"):
        note = (note + "; " if note else "") + "std n/a"
    if np.isnan(mean):
        note = (note + "; " if note else "") + f"could not read {text!r}"
    return mean, std, note


def read_table_rows(
    xlsx_path: Path, *, sheet: str = "Cloud&MetQuantities", author_contains: str = "Kavin"
) -> pd.DataFrame:
    """Rows of the seasonal-averages sheet whose Author(s) column (J) contains `author_contains`.

    One row per spreadsheet row, indexed by the spreadsheet row number, with
    the descriptive columns (quantity, units, location, method, author, note)
    and, for each season, <season>_raw (the cell text), <season>_mean,
    <season>_std and <season>_note."""
    from openpyxl import load_workbook  # read-only; pandas would coerce the text cells

    wb = load_workbook(Path(xlsx_path).expanduser(), read_only=True, data_only=False)
    ws = wb[sheet]
    records = {}
    for row in ws.iter_rows(min_row=3):
        cells = {c.column_letter: c.value for c in row if c.value is not None and hasattr(c, "column_letter")}
        author = str(cells.get("J", ""))
        if author_contains.lower() not in author.lower():
            continue
        r = row[0].row
        rec = {
            "quantity": str(cells.get("A", "")).strip(),
            "units": str(cells.get("B", "") or "").strip(),
            "location": str(cells.get("C", "") or "").strip(),
            "method": str(cells.get("D", "") or "").strip(),
            "author": author.strip(),
            "note": str(cells.get("K", "") or "").strip(),
        }
        for season, col in SEASON_COLUMNS.items():
            raw = cells.get(col)
            mean, std, note = parse_cell(raw)
            rec.update({f"{season}_raw": raw, f"{season}_mean": mean, f"{season}_std": std, f"{season}_note": note})
        records[r] = rec
    wb.close()
    return pd.DataFrame.from_dict(records, orient="index").rename_axis("sheet_row")


def table_values(table: pd.DataFrame, sheet_row: int) -> pd.DataFrame:
    """Kavin's mean and std per season for one spreadsheet row (rows = SEASON_ORDER)."""
    rec = table.loc[sheet_row]
    return pd.DataFrame(
        {
            "mean": [rec[f"{s}_mean"] for s in SEASON_ORDER],
            "std": [rec[f"{s}_std"] for s in SEASON_ORDER],
            "raw": [rec[f"{s}_raw"] for s in SEASON_ORDER],
        },
        index=SEASON_ORDER,
    )


# ---------------------------------------------------------------------------
# side-by-side comparison
# ---------------------------------------------------------------------------
def format_pm(mean: float, std: float, sig: int = 2) -> str:
    """'mean ± std' with `sig` significant figures on the std (and mean to the same decimal)."""
    if not np.isfinite(mean):
        return "–"
    if not np.isfinite(std) or std == 0:
        # no spread given: three significant figures, never scientific notation
        return np.format_float_positional(mean, precision=3, unique=False, fractional=False, trim="-")
    decimals = max(0, sig - 1 - int(np.floor(np.log10(abs(std)))))
    return f"{mean:.{decimals}f} ± {std:.{decimals}f}"


def compare(
    ours: pd.DataFrame,
    table: pd.DataFrame,
    sheet_row: int,
    definition: str,
    *,
    close_pct: float = 10.0,
) -> pd.DataFrame:
    """Our seasonal statistics next to the spreadsheet's, one row per season.

    Columns: sheet_row, quantity, definition, season, ours, n, table,
    diff_pct = 100 (ours - table) / table on the means, std_ratio = our std /
    table std, and close = |diff_pct| <= close_pct. `close_pct` is only a
    reading aid for scanning the summary, not a statistical test."""
    return compare_to(
        ours, table_values(table, sheet_row), sheet_row, table.loc[sheet_row, "quantity"], definition,
        close_pct=close_pct,
    )


def compare_to(
    ours: pd.DataFrame,
    kav: pd.DataFrame,
    sheet_row,
    quantity: str,
    definition: str,
    *,
    close_pct: float = 10.0,
) -> pd.DataFrame:
    """As ``compare``, against any reference with columns mean, std, raw (rows = SEASON_ORDER).

    Used for references derived from several sheet rows, e.g. the
    hour-weighted combination of the cloud and haze residual rows."""
    rows = []
    for season in SEASON_ORDER:
        m, s, n = ours.loc[season, "mean"], ours.loc[season, "std"], ours.loc[season, "n"]
        km, ks = kav.loc[season, "mean"], kav.loc[season, "std"]
        diff = 100.0 * (m - km) / km if np.isfinite(km) and km != 0 and np.isfinite(m) else np.nan
        rows.append(
            {
                "sheet_row": sheet_row,
                "quantity": quantity,
                "definition": definition,
                "season": season,
                "ours_mean": m,
                "ours_std": s,
                "n": int(n),
                "table_mean": km,
                "table_std": ks,
                "ours": format_pm(m, s),
                "table": kav.loc[season, "raw"],
                "diff_pct": diff,
                "std_ratio": s / ks if np.isfinite(ks) and ks > 0 else np.nan,
                "close": bool(np.isfinite(diff) and abs(diff) <= close_pct),
            }
        )
    return pd.DataFrame(rows)


def show(comparison: pd.DataFrame) -> pd.DataFrame:
    """The columns worth reading in a notebook display."""
    cols = ["season", "ours", "n", "table", "diff_pct", "std_ratio", "close"]
    out = comparison[cols].copy()
    out["diff_pct"] = out["diff_pct"].round(1)
    out["std_ratio"] = out["std_ratio"].round(2)
    return out.set_index("season")


def closest_definition(comparisons: Sequence[pd.DataFrame]) -> pd.DataFrame:
    """Rank candidate definitions by how well they reproduce the table.

    Used where Kavin's definition is not documented (e.g. the rain rows).
    Ranked by the median |diff_pct| of the means over the five columns (the
    median, because a percentage of a near-zero table value, e.g. spring
    rain, can dominate a mean). median_abs_log2_std_ratio says how well the
    spread matches as well (0 = identical std, 1 = a factor of 2). The best
    candidate is a hypothesis about Kavin's method, not proof of it."""
    rows = []
    for c in comparisons:
        d = c["diff_pct"].abs()
        r = np.log2(c["std_ratio"].astype(float)).abs()
        rows.append(
            {
                "definition": c["definition"].iloc[0],
                "median_abs_diff_pct": d.median(),
                "max_abs_diff_pct": d.max(),
                "median_abs_log2_std_ratio": r.median(),
                "seasons_compared": int(d.notna().sum()),
            }
        )
    return pd.DataFrame(rows).sort_values("median_abs_diff_pct").reset_index(drop=True)
